"""Офлайн-проверки двух независимых Excel-бланков пробной разметки H2."""

from __future__ import annotations

import hashlib
import json
import tempfile
import unittest

from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

from src.corpus.h2_practice import (
    PreparedH2Practice,
    prepare_h2_practice,
    publish_h2_practice,
)

CARD_CELLS = {
    "annotation_item_id": "A3",
    "input_sha256": "A4",
    "title": "A7",
    "abstract": "A10",
    "familiarity": "B13",
    "decision": "B14",
    "confidence": "B15",
    "rationale": "B16",
    "completed_at": "B17",
}
WORKBOOK_NAMES = ("Разметка_Дмитрий.xlsx", "Разметка_научрук.xlsx")
ANSWER_FILES = {"answers.jsonl", "comparison.jsonl", "disagreements.jsonl"}


def _json_bytes(value: object) -> bytes:
    """Сериализовать синтетические данные в устойчивый UTF-8 JSON."""

    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")


def _fingerprint(payload: bytes) -> dict[str, object]:
    """Получить размер и контрольную сумму точных байтов тестового файла."""

    return {"bytes": len(payload), "sha256": hashlib.sha256(payload).hexdigest()}


class H2PracticeTests(unittest.TestCase):
    """Проверить безопасный приём ответов без настоящих статей и рабочих реестров."""

    def setUp(self) -> None:
        """Подготовить две синтетические карточки, шаблоны и заполненные копии."""

        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name).resolve()
        self.packet_dir = self.root / "manifests/results/practice-packet"
        self.template_dir = self.root / "outputs/practice-templates"
        self.answer_dir = self.root / "outputs/filled-answers"
        self.output_dir = self.root / "manifests/results/practice-check-v1"

        for directory in (self.packet_dir / "author", self.template_dir, self.answer_dir):
            directory.mkdir(parents=True)

        self.items = []

        for item_index in range(1, 3):
            text = {
                "title": f"Исследование физических процессов {item_index}",
                "abstract": f"Изучено движение частиц {item_index}.\nПриведены результаты измерений.",
            }
            self.items.append({
                "annotation_item_id": f"practice-item-{item_index}",
                **text,
                "input_sha256": hashlib.sha256(_json_bytes(text)).hexdigest(),
            })

        inputs = b"".join(_json_bytes(item) + b"\n" for item in self.items)
        instructions = "Черновик памятки. Разметка независимая. Памятка ещё не утверждена.\n".encode()
        (self.packet_dir / "inputs.jsonl").write_bytes(inputs)
        (self.packet_dir / "author/INSTRUCTIONS.txt").write_bytes(instructions)
        self.packet = {
            "format_version": "h2-practice-draft-v1",
            "packet_id": "synthetic-practice",
            "guide_version": "synthetic-guide-v1",
            "status": "awaiting_guide_approval",
            "guide_approved": False,
            "annotation_started": False,
            "formal_calibration": False,
            "formal_audit": False,
            "training_ready": False,
            "item_count": len(self.items),
            "output_files": {
                "inputs.jsonl": _fingerprint(inputs),
                "author/INSTRUCTIONS.txt": _fingerprint(instructions),
            },
        }
        packet_bytes = _json_bytes(self.packet)
        (self.packet_dir / "packet.json").write_bytes(packet_bytes)
        workbook_bytes = self._make_template()
        self.verification = {
            "format_version": "h2-practice-excel-v1",
            "source_packet_sha256": hashlib.sha256(packet_bytes).hexdigest(),
            "source_inputs_sha256": hashlib.sha256(inputs).hexdigest(),
            "source_guide_sha256": hashlib.sha256(instructions).hexdigest(),
            "files": list(WORKBOOK_NAMES),
            "initial_workbook_sha256": hashlib.sha256(workbook_bytes).hexdigest(),
            "items": len(self.items),
            "card_cells": CARD_CELLS,
            "participant_cell": "Памятка!B4",
            "sheet_order": ["Памятка", "01", "02"],
        }
        (self.template_dir / "verification.json").write_bytes(_json_bytes(self.verification))
        self.author_file = self.answer_dir / "author.xlsx"
        self.supervisor_file = self.answer_dir / "supervisor.xlsx"

        for path, participant in (
            (self.author_file, "Участник первый"),
            (self.supervisor_file, "Участник второй"),
        ):
            path.write_bytes(workbook_bytes)
            self._fill_answers(path, participant)

        registry = self.root / "manifests/works.jsonl"
        registry.write_bytes(b'{"work_id":"untouched-synthetic-work"}\n')

    def _make_template(self) -> bytes:
        """Создать один пустой шаблон и сохранить две побайтово одинаковые копии."""

        workbook = Workbook()
        guide = workbook.worksheets[0]
        guide.title = "Памятка"
        guide["A1"] = "Черновик: памятка не утверждена"
        guide["A6"] = "Отвечайте независимо и только по показанному тексту."

        for item_index, item in enumerate(self.items, start=1):
            sheet = workbook.create_sheet(f"{item_index:02d}")
            sheet["A1"] = f"Карточка {item_index}"
            sheet["A2"] = "Сначала согласуйте памятку"

            for field in ("annotation_item_id", "input_sha256", "title", "abstract"):
                sheet[CARD_CELLS[field]] = item[field]

        workbook.save(self.template_dir / WORKBOOK_NAMES[0])
        workbook.close()
        payload = (self.template_dir / WORKBOOK_NAMES[0]).read_bytes()
        (self.template_dir / WORKBOOK_NAMES[1]).write_bytes(payload)
        return payload

    def _fill_answers(self, path: Path, participant: str) -> None:
        """Заполнить только разрешённые поля синтетического Excel-бланка."""

        workbook = load_workbook(path)
        workbook["Памятка"]["B4"] = participant

        for sheet_name in ("01", "02"):
            sheet = workbook[sheet_name]
            sheet["B13"] = "нет"
            sheet["B14"] = "C1"
            sheet["B15"] = "high"
            sheet["B16"] = "В аннотации прямо указана исследуемая физическая задача."
            sheet["B17"] = "2026-10-11T14:30:00+03:00"

        workbook.save(path)
        workbook.close()

    def _edit_cell(self, path: Path, sheet_name: str, address: str, value: Any) -> None:
        """Изменить одну ячейку тестовой копии, не затрагивая исходный шаблон."""

        workbook = load_workbook(path)
        workbook[sheet_name][address] = value
        workbook.save(path)
        workbook.close()

    def _prepare(self, **overrides: Any) -> PreparedH2Practice:
        """Вызвать проверку с локальными путями и необязательными заменами аргументов."""

        arguments = {
            "project_root": self.root,
            "packet_dir": self.packet_dir,
            "template_dir": self.template_dir,
            "author_file": self.author_file,
            "supervisor_file": self.supervisor_file,
            "output_dir": self.output_dir,
        }
        arguments.update(overrides)
        return prepare_h2_practice(**arguments)

    def _tree_bytes(self) -> dict[str, bytes]:
        """Снять содержимое всех обычных файлов для проверки отсутствия записи."""

        return {
            path.relative_to(self.root).as_posix(): path.read_bytes()
            for path in self.root.rglob("*")
            if path.is_file() and not path.is_symlink()
        }

    def _replace_packet_inputs(self, items: list[dict[str, str]]) -> None:
        """Обновить контрольные суммы тестового пакета, но не исправлять его смысл."""

        payload = b"".join(_json_bytes(item) + b"\n" for item in items)
        (self.packet_dir / "inputs.jsonl").write_bytes(payload)
        description = json.loads((self.packet_dir / "packet.json").read_bytes())
        description["output_files"]["inputs.jsonl"] = _fingerprint(payload)
        packet_bytes = _json_bytes(description)
        (self.packet_dir / "packet.json").write_bytes(packet_bytes)
        verification = json.loads((self.template_dir / "verification.json").read_bytes())
        verification["source_packet_sha256"] = hashlib.sha256(packet_bytes).hexdigest()
        verification["source_inputs_sha256"] = hashlib.sha256(payload).hexdigest()
        (self.template_dir / "verification.json").write_bytes(_json_bytes(verification))

    def _rows(self, prepared: PreparedH2Practice, filename: str) -> list[dict[str, Any]]:
        """Прочитать подготовленные JSONL-строки без публикации на диск."""

        return [json.loads(line) for line in prepared.files[filename].splitlines()]

    def _assert_invalid(self, prepared: PreparedH2Practice) -> None:
        """Проверить, что ошибочный ввод не выпускается как принятая разметка."""

        self.assertEqual(prepared.report["status"], "invalid")
        self.assertGreater(prepared.report["counts"]["errors"], 0)
        self.assertFalse(ANSWER_FILES.intersection(prepared.files))
        self.assertFalse(prepared.report["formal_h2_ready"])

    def test_complete_agreement_is_checked_without_approving_protocol(self) -> None:
        """Совпавшие ответы остаются пробной разметкой, а памятка не утверждается."""

        before = self._tree_bytes()
        prepared = self._prepare()

        self.assertEqual(self._tree_bytes(), before)
        self.assertFalse(self.output_dir.exists())
        self.assertEqual(prepared.report["status"], "practice_answers_checked")
        self.assertEqual(prepared.report["counts"]["items"], 2)
        self.assertEqual(prepared.report["counts"]["errors"], 0)
        self.assertEqual(prepared.report["counts"]["complete_pairs"], 2)
        self.assertEqual(prepared.report["counts"]["disagreements"], 0)
        self.assertFalse(prepared.report["formal_h2_ready"])
        self.assertFalse(prepared.report["guide_approved"])
        self.assertEqual(len(self._rows(prepared, "answers.jsonl")), 4)
        self.assertEqual(self._rows(prepared, "disagreements.jsonl"), [])
        comparisons = self._rows(prepared, "comparison.jsonl")
        self.assertEqual([row["annotation_item_id"] for row in comparisons], [item["annotation_item_id"] for item in self.items])
        self.assertTrue(all(row["decision_equal"] for row in comparisons))

    def test_disagreement_preserves_both_answers_without_final_label(self) -> None:
        """Разногласие попадает в отдельный файл, итоговый класс не назначается."""

        self._edit_cell(self.supervisor_file, "02", "B14", "AMBIGUOUS")
        prepared = self._prepare()
        disagreements = self._rows(prepared, "disagreements.jsonl")

        self.assertEqual(prepared.report["counts"]["disagreements"], 1)
        self.assertEqual(len(disagreements), 1)
        self.assertEqual(disagreements[0]["annotation_item_id"], self.items[1]["annotation_item_id"])
        self.assertEqual(disagreements[0]["author_decision"], "C1")
        self.assertEqual(disagreements[0]["supervisor_decision"], "AMBIGUOUS")
        self.assertFalse(disagreements[0]["decision_equal"])

        for filename in ANSWER_FILES:
            for row in self._rows(prepared, filename):
                self.assertNotIn("label", row)
                self.assertNotIn("final_label", row)

    def test_output_keeps_exact_answer_fields_and_source_fingerprints(self) -> None:
        """Ответы сохраняют роль, текст объяснения, время и контрольную сумму карточки."""

        prepared = self._prepare()
        rows = self._rows(prepared, "answers.jsonl")

        for row in rows:
            self.assertIn(row["role"], {"author", "supervisor"})
            self.assertIn(row["annotator_id"], {"Участник первый", "Участник второй"})
            self.assertEqual(row["decision"], "C1")
            self.assertEqual(row["familiarity"], "нет")
            self.assertEqual(row["confidence"], "high")
            self.assertEqual(row["completed_at"], "2026-10-11T14:30:00+03:00")
            self.assertEqual(row["rationale"], "В аннотации прямо указана исследуемая физическая задача.")
            item = next(item for item in self.items if item["annotation_item_id"] == row["annotation_item_id"])
            self.assertEqual(row["input_sha256"], item["input_sha256"])

        for relative, digest in prepared.source_hashes.items():
            self.assertEqual(hashlib.sha256((self.root / relative).read_bytes()).hexdigest(), digest)

    def test_blank_workbooks_are_invalid_not_complete_agreement(self) -> None:
        """Два пустых бланка не превращаются в нулевое число ошибок и полное согласие."""

        for path in (self.author_file, self.supervisor_file):
            path.write_bytes((self.template_dir / WORKBOOK_NAMES[0]).read_bytes())

        prepared = self._prepare()
        self._assert_invalid(prepared)
        self.assertEqual(prepared.report["counts"]["complete_pairs"], 0)

    def test_missing_required_answer_field_blocks_answer_export(self) -> None:
        """Любой незаполненный обязательный ответ запрещает выпуск принятых строк."""

        original = self.author_file.read_bytes()

        for address in ("B13", "B14", "B15", "B16", "B17"):
            with self.subTest(address=address):
                self.author_file.write_bytes(original)
                self._edit_cell(self.author_file, "01", address, None)
                self._assert_invalid(self._prepare())

    def test_invalid_enums_and_whitespace_only_rationale_are_rejected(self) -> None:
        """Вставленные мимо выпадающих списков значения проверяются повторно."""

        original = self.author_file.read_bytes()

        for address, value in (("B13", "maybe"), ("B14", "C6"), ("B15", "sure"), ("B16", "   \n")):
            with self.subTest(address=address, value=value):
                self.author_file.write_bytes(original)
                self._edit_cell(self.author_file, "01", address, value)
                self._assert_invalid(self._prepare())

    def test_completion_timestamp_requires_real_date_and_timezone(self) -> None:
        """Локальное время, дата без времени и невозможный день не принимаются."""

        original = self.author_file.read_bytes()

        for timestamp in (
            "2026-10-11", "2026-10-11T14:30:00", "2026-02-30T14:30:00+03:00", "вчера",
            "2026-10-11T14:30:00+00:60", "2026-10-11T14:30:00+03:99", "2026-10-11T14:30:00-00:60",
        ):
            with self.subTest(timestamp=timestamp):
                self.author_file.write_bytes(original)
                self._edit_cell(self.author_file, "01", "B17", timestamp)
                self._assert_invalid(self._prepare())

    def test_same_normalized_person_is_not_independent_annotation(self) -> None:
        """Регистр и внешние пробелы не превращают одного участника в двух."""

        self._edit_cell(self.supervisor_file, "Памятка", "B4", "  УЧАСТНИК ПЕРВЫЙ  ")
        self._assert_invalid(self._prepare())

    def test_same_input_path_is_rejected_before_writes(self) -> None:
        """Один файл нельзя передать за два независимых бланка."""

        before = self._tree_bytes()

        with self.assertRaises(ValueError):
            self._prepare(supervisor_file=self.author_file)

        self.assertEqual(self._tree_bytes(), before)

    def test_hardlinks_cannot_represent_two_independent_answers(self) -> None:
        """Разные имена одного физического файла не считаются независимыми ответами."""

        linked_input = self.answer_dir / "hardlinked-author.xlsx"
        linked_input.hardlink_to(self.author_file)
        before = self._tree_bytes()

        with self.assertRaises(ValueError):
            self._prepare(supervisor_file=linked_input)

        self.assertEqual(self._tree_bytes(), before)

    def test_changed_identity_text_hash_and_guide_are_detected(self) -> None:
        """Правки защищённых данных и памятки обнаруживаются независимо от ответов."""

        original = self.author_file.read_bytes()

        for sheet_name, address, value in (
            ("01", "A3", "different-id"),
            ("01", "A4", "a" * 64),
            ("01", "A7", "Другое название"),
            ("01", "A10", "Другая аннотация"),
            ("Памятка", "A6", "Обсудите ответы заранее"),
        ):
            with self.subTest(sheet=sheet_name, address=address):
                self.author_file.write_bytes(original)
                self._edit_cell(self.author_file, sheet_name, address, value)
                self._assert_invalid(self._prepare())

    def test_formula_is_not_evaluated_or_exported(self) -> None:
        """Формула вместо ответа отклоняется без вычисления и без записи результата."""

        self._edit_cell(self.author_file, "01", "B16", '=HYPERLINK("https://example.invalid","текст")')
        before = self._tree_bytes()

        try:
            prepared = self._prepare()

        except ValueError:
            prepared = None

        if prepared is not None:
            self._assert_invalid(prepared)

        self.assertEqual(self._tree_bytes(), before)

    def test_sheet_reordering_and_unexpected_cell_are_not_accepted(self) -> None:
        """Порядок листов и новые ячейки вне полей ответа также проверяются."""

        original = self.author_file.read_bytes()
        workbook = load_workbook(self.author_file)
        workbook.move_sheet("02", offset=-1)
        workbook.save(self.author_file)
        workbook.close()
        self._assert_invalid(self._prepare())
        self.author_file.write_bytes(original)
        self._edit_cell(self.author_file, "01", "C18", "Дополнительная подсказка из внешнего источника")
        self._assert_invalid(self._prepare())

    def test_corrupted_packet_digest_is_rejected(self) -> None:
        """Исходные карточки проверяются по точным байтам, а не только по структуре."""

        path = self.packet_dir / "inputs.jsonl"
        path.write_bytes(path.read_bytes() + b"\n")

        with self.assertRaises(ValueError):
            self._prepare()

        self.assertFalse(self.output_dir.exists())

    def test_changed_template_and_invalid_receipt_are_rejected(self) -> None:
        """Подмена пустого шаблона или схемы ячеек не меняет допустимый протокол."""

        template = self.template_dir / WORKBOOK_NAMES[0]
        original = template.read_bytes()
        self._edit_cell(template, "01", "A7", "Подменённый шаблон")

        with self.assertRaises(ValueError):
            self._prepare()

        template.write_bytes(original)
        self.verification["card_cells"] = {**CARD_CELLS, "decision": "Z99"}
        (self.template_dir / "verification.json").write_bytes(_json_bytes(self.verification))

        with self.assertRaises(ValueError):
            self._prepare()

    def test_duplicate_ids_and_false_input_hash_are_rejected(self) -> None:
        """Верная опись файлов не скрывает повтор ID или неверный хеш текста."""

        for mutation in ("duplicate_id", "false_hash"):
            with self.subTest(mutation=mutation):
                items = [dict(item) for item in self.items]

                if mutation == "duplicate_id":
                    items[1]["annotation_item_id"] = items[0]["annotation_item_id"]

                else:
                    items[0]["input_sha256"] = "a" * 64

                self._replace_packet_inputs(items)

                with self.assertRaises(ValueError):
                    self._prepare()

    def test_template_must_match_packet_even_with_matching_file_hashes(self) -> None:
        """Шаблон не становится источником истины при подмене текста и его хеша."""

        path = self.template_dir / WORKBOOK_NAMES[0]
        self._edit_cell(path, "01", "A7", "Другое название в обоих шаблонах")
        payload = path.read_bytes()
        (self.template_dir / WORKBOOK_NAMES[1]).write_bytes(payload)
        self.verification["initial_workbook_sha256"] = hashlib.sha256(payload).hexdigest()
        (self.template_dir / "verification.json").write_bytes(_json_bytes(self.verification))

        with self.assertRaises(ValueError):
            self._prepare()

    def test_publication_preserves_all_sources_and_registry(self) -> None:
        """Публикация добавляет только новый пакет и не изменяет ответы и реестр."""

        before = self._tree_bytes()
        prepared = self._prepare()

        self.assertTrue(publish_h2_practice(prepared))
        self.assertEqual({path.name for path in self.output_dir.iterdir()}, set(prepared.files))

        for relative, payload in before.items():
            self.assertEqual((self.root / relative).read_bytes(), payload)

        for name, payload in prepared.files.items():
            self.assertEqual((self.output_dir / name).read_bytes(), payload)

    def test_dry_run_never_creates_output(self) -> None:
        """Пробная публикация оставляет всё дерево проекта без изменений."""

        before = self._tree_bytes()
        publish_h2_practice(self._prepare(), dry_run=True)

        self.assertEqual(self._tree_bytes(), before)
        self.assertFalse(self.output_dir.exists())

    def test_existing_output_is_never_overwritten_even_on_exact_replay(self) -> None:
        """Повтор требует новый каталог и не затрагивает уже сохранённый результат."""

        prepared = self._prepare()
        publish_h2_practice(prepared)
        before = self._tree_bytes()

        with self.assertRaises(ValueError):
            publish_h2_practice(prepared)

        self.assertEqual(self._tree_bytes(), before)

    def test_source_change_after_preparation_stops_publication(self) -> None:
        """Правка ответа между проверкой и записью обнаруживается перед публикацией."""

        prepared = self._prepare()
        self._edit_cell(self.author_file, "01", "B16", "Новое объяснение участника.")
        before = self._tree_bytes()

        with self.assertRaises(ValueError):
            publish_h2_practice(prepared)

        self.assertEqual(self._tree_bytes(), before)
        self.assertFalse(self.output_dir.exists())

    def test_prepared_payload_and_report_cannot_change_silently(self) -> None:
        """Редактирование подготовленных в памяти байтов или отчёта требует новой проверки."""

        before = self._tree_bytes()

        for mutation in ("payload", "report"):
            with self.subTest(mutation=mutation):
                prepared = self._prepare()

                if mutation == "payload":
                    prepared.files["answers.jsonl"] += b"\n"

                else:
                    prepared.report["formal_h2_ready"] = True

                with self.assertRaises(ValueError):
                    publish_h2_practice(prepared)

        self.assertEqual(self._tree_bytes(), before)
        self.assertFalse(self.output_dir.exists())

    def test_source_replaced_by_symlink_after_preparation_is_rejected(self) -> None:
        """После подготовки нельзя подменить вход ссылкой даже с теми же байтами."""

        prepared = self._prepare()
        relocated = self.answer_dir / "author-relocated.xlsx"
        self.author_file.rename(relocated)
        self.author_file.symlink_to(relocated)
        before = self._tree_bytes()

        with self.assertRaises(ValueError):
            publish_h2_practice(prepared)

        self.assertEqual(self._tree_bytes(), before)
        self.assertTrue(self.author_file.is_symlink())
        self.assertFalse(self.output_dir.exists())

    def test_unsafe_output_paths_are_rejected(self) -> None:
        """Результат нельзя направить в источники, корень или за пределы проекта."""

        before = self._tree_bytes()

        for output_dir in (
            self.root,
            self.packet_dir,
            self.packet_dir / "new-output",
            self.template_dir / "new-output",
            self.root / "outputs/new-output",
            self.root.parent / f"{self.root.name}-outside",
        ):
            with self.subTest(output_dir=output_dir), self.assertRaises(ValueError):
                self._prepare(output_dir=output_dir)

        self.assertEqual(self._tree_bytes(), before)

    def test_symlink_input_and_output_are_rejected(self) -> None:
        """Символические ссылки не обходят проверку путей даже внутри проекта."""

        linked_input = self.answer_dir / "linked-author.xlsx"
        linked_input.symlink_to(self.author_file)

        with self.assertRaises(ValueError):
            self._prepare(author_file=linked_input)

        linked_parent = self.root / "manifests/results/linked-output"
        linked_parent.symlink_to(self.answer_dir, target_is_directory=True)

        with self.assertRaises(ValueError):
            self._prepare(output_dir=linked_parent / "result")


if __name__ == "__main__":
    unittest.main()
