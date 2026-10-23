"""Проверки безопасного чтения синтетических Excel-бланков разметки H2."""

from __future__ import annotations

import io
import unittest
import warnings
import zipfile

from datetime import datetime
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from src.corpus.h2_practice_xlsx import PracticeWorkbook, read_practice_workbook


def _workbook_bytes(workbook: Workbook) -> bytes:
    """Сохранить только синтетическую книгу в памяти, не касаясь файлов пользователя."""

    destination = io.BytesIO()
    workbook.save(destination)
    workbook.close()

    return destination.getvalue()


def _replace_parts(data: bytes, replacements: dict[str, bytes]) -> bytes:
    """Изменить тестовые XML-части архива либо добавить намеренно недопустимую часть."""

    destination = io.BytesIO()

    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(destination, "w") as target:
        for entry in source.infolist():
            target.writestr(entry.filename, replacements.get(entry.filename, source.read(entry)))

        for name, payload in replacements.items():
            if name not in source.namelist():
                target.writestr(name, payload)

    return destination.getvalue()


class PracticeXlsxTests(unittest.TestCase):
    """Бланк читается без нормализации текста, вычислений и скрытых внешних связей."""

    def setUp(self) -> None:
        """Подготовить небольшую книгу с памяткой, карточкой и выпадающим списком."""

        self.workbook = Workbook()
        self.instructions = self.workbook.worksheets[0]
        self.instructions.title = "Памятка"
        self.instructions["A1"] = "Инструкция"
        self.card = self.workbook.create_sheet("01")
        self.card["A3"] = "annotation_item_id: practice-01"
        self.card["A7"] = "  Точное название\nс пробелами  "
        self.card["A10"] = "Исходная аннотация: ё, е́, α."
        self.card.merge_cells("A10:D10")
        self.card["B14"] = "C1"
        self.card.print_area = "A1:D19"
        validation = DataValidation(type="list", formula1='"C1,C2,OTHER"')
        self.card.add_data_validation(validation)
        validation.add("B14")
        self.workbook.create_sheet("Пустой лист")

    def test_reads_exact_strings_and_sheet_order(self) -> None:
        """Объединённые ячейки и списки допустимы; пробелы и Unicode не меняются."""

        data = _workbook_bytes(self.workbook)
        original = bytes(data)
        result = read_practice_workbook(data)

        self.assertIsInstance(result, PracticeWorkbook)
        self.assertEqual(result.sheets, ("Памятка", "01", "Пустой лист"))
        self.assertEqual(result.cells["01"]["A7"], "  Точное название\nс пробелами  ")
        self.assertEqual(result.cells["01"]["A10"], "Исходная аннотация: ё, е́, α.")
        self.assertEqual(result.cells["01"]["B14"], "C1")
        self.assertNotIn("B13", result.cells["01"])
        self.assertEqual(result.cells["Пустой лист"], {})
        self.assertEqual(data, original)

    def test_rejects_nontext_values(self) -> None:
        """Числа, логические значения и даты не превращаются в строки молча."""

        for value in (123, 0.5, True, datetime(2026, 10, 11, 12)):
            with self.subTest(value=value):
                self.card["B17"] = value

                with self.assertRaisesRegex(ValueError, "ожидается текст"):
                    read_practice_workbook(_workbook_bytes(self.workbook))

    def test_rejects_formulas_and_error_cells(self) -> None:
        """Формулы и ошибки Excel запрещены, даже если сохранён вычисленный результат."""

        for value in ("=1+1", "#DIV/0!"):
            with self.subTest(value=value):
                self.card["B14"] = value

                with self.assertRaisesRegex(ValueError, "формулы"):
                    read_practice_workbook(_workbook_bytes(self.workbook))

    def test_rejects_hidden_sheets(self) -> None:
        """Ни обычный, ни усиленный режим сокрытия листов не допускается."""

        for state in ("hidden", "veryHidden"):
            with self.subTest(state=state):
                self.card.sheet_state = state

                with self.assertRaisesRegex(ValueError, "скрытых листов"):
                    read_practice_workbook(_workbook_bytes(self.workbook))

    def test_rejects_comments(self) -> None:
        """Комментарии могут содержать чужие ответы и не принимаются вместе с бланком."""

        self.card["B14"].comment = Comment("Чужое решение", "Автор")

        with self.assertRaises(ValueError):
            read_practice_workbook(_workbook_bytes(self.workbook))

    def test_rejects_external_and_internal_hyperlinks(self) -> None:
        """Ссылки не открываются и не включаются в распознанные данные."""

        for target in ("https://example.org/", "#Памятка!A1"):
            with self.subTest(target=target):
                self.card["B14"].hyperlink = target

                with self.assertRaises(ValueError):
                    read_practice_workbook(_workbook_bytes(self.workbook))

    def test_rejects_out_of_bounds_sheet(self) -> None:
        """Даже одна ячейка за разрешённой областью приводит к отказу чтения."""

        for coordinate in ("A201", "I1"):
            with self.subTest(coordinate=coordinate):
                self.card[coordinate] = "лишняя ячейка"

                with self.assertRaisesRegex(ValueError, "Размер листа"):
                    read_practice_workbook(_workbook_bytes(self.workbook))

                del self.card[coordinate]

    def test_rejects_too_many_sheets(self) -> None:
        """Количество листов ограничивается до загрузки книги библиотекой."""

        for index in range(30):
            self.workbook.create_sheet(f"Дополнительный {index}")

        with self.assertRaisesRegex(ValueError, "не больше 32"):
            read_practice_workbook(_workbook_bytes(self.workbook))

    def test_does_not_trust_declared_small_dimensions(self) -> None:
        """Ячейки не теряются, если редактор указал слишком маленькую область листа."""

        data = _workbook_bytes(self.workbook)

        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            sheet = archive.read("xl/worksheets/sheet2.xml")

        changed = _replace_parts(data, {
            "xl/worksheets/sheet2.xml": sheet.replace(b'ref="A3:D14"', b'ref="A1:A1"'),
        })

        self.assertEqual(read_practice_workbook(changed).cells["01"]["B14"], "C1")

    def test_rejects_invalid_archives_and_xml(self) -> None:
        """Повреждения контейнера и XML сообщаются как понятная ошибка входных данных."""

        data = _workbook_bytes(self.workbook)
        damaged = _replace_parts(data, {"xl/workbook.xml": b"<broken>"})

        for payload in (b"not a zip", damaged):
            with self.subTest(payload=payload[:10]), self.assertRaisesRegex(ValueError, "Excel-бланк"):
                read_practice_workbook(payload)

    def test_rejects_empty_input(self) -> None:
        """Пустой ввод не выдаётся за бланк без заполненных ответов."""

        with self.assertRaisesRegex(ValueError, "непустые байты"):
            read_practice_workbook(b"")

    def test_rejects_external_relationship_without_hyperlink(self) -> None:
        """Произвольная внешняя связь запрещена независимо от типа и наличия ячейки."""

        data = _workbook_bytes(self.workbook)
        relationships = (
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="custom" Target="https://example.org/" TargetMode="External"/>'
            '</Relationships>'
        )

        with self.assertRaisesRegex(ValueError, "внешние связи"):
            read_practice_workbook(_replace_parts(data, {
                "xl/worksheets/_rels/sheet2.xml.rels": relationships.encode("utf-8"),
            }))

    def test_rejects_named_expressions(self) -> None:
        """Дополнительные имена с выражениями запрещены, области печати допустимы."""

        self.workbook.defined_names.add(DefinedName("HiddenFormula", attr_text="1+1"))

        with self.assertRaisesRegex(ValueError, "именованные выражения"):
            read_practice_workbook(_workbook_bytes(self.workbook))

    def test_rejects_duplicate_coordinates(self) -> None:
        """Повторяющаяся координата не позволяет незаметно подменить ответ."""

        data = _workbook_bytes(self.workbook)

        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            sheet = archive.read("xl/worksheets/sheet2.xml")

        changed = sheet.replace(
            b'<c r="B14" t="inlineStr"><is><t>C1</t></is></c>',
            b'<c r="B14" t="inlineStr"><is><t>C1</t></is></c>'
            b'<c r="B14" t="inlineStr"><is><t>C2</t></is></c>',
        )
        self.assertNotEqual(sheet, changed)

        with self.assertRaisesRegex(ValueError, "повторяющаяся координата"):
            read_practice_workbook(_replace_parts(data, {"xl/worksheets/sheet2.xml": changed}))

    def test_rejects_dtd_and_entity_declarations(self) -> None:
        """Декларации сущностей отклоняются до XML-разбора, включая UTF-16."""

        data = _workbook_bytes(self.workbook)
        text = '<!DOCTYPE workbook [<!ENTITY secret "value">]><workbook/>'

        for encoding in ("utf-8", "utf-16", "utf-32"):
            with self.subTest(encoding=encoding), self.assertRaisesRegex(ValueError, "DTD"):
                read_practice_workbook(_replace_parts(data, {"xl/workbook.xml": text.encode(encoding)}))

    def test_rejects_suspicious_or_active_parts(self) -> None:
        """Макросы, OLE, внешние книги и обход каталогов недопустимы в текстовом бланке."""

        data = _workbook_bytes(self.workbook)

        for name in (
            "xl/vbaProject.bin", "xl/embeddings/oleObject1.bin", "xl/externalLinks/externalLink1.xml",
            "../xl/workbook.xml", "/xl/workbook.xml", "xl\\workbook.xml",
        ):
            with self.subTest(name=name), self.assertRaises(ValueError):
                read_practice_workbook(_replace_parts(data, {name: b"<extra/>"}))

    def test_rejects_duplicate_parts(self) -> None:
        """Повторяющиеся ZIP-имена не дают библиотеке выбрать другую версию XML."""

        destination = io.BytesIO(_workbook_bytes(self.workbook))

        with warnings.catch_warnings(), zipfile.ZipFile(destination, "a") as archive:
            warnings.simplefilter("ignore", UserWarning)
            archive.writestr("xl/workbook.xml", b"<workbook/>")

        with self.assertRaisesRegex(ValueError, "повторяющаяся часть"):
            read_practice_workbook(destination.getvalue())

    def test_rejects_oversized_inputs_and_archives(self) -> None:
        """Пределы файла, распакованных данных, частей и их количества независимы."""

        data = _workbook_bytes(self.workbook)

        for constant in ("MAX_INPUT_BYTES", "MAX_UNCOMPRESSED_BYTES", "MAX_PART_BYTES", "MAX_ARCHIVE_ENTRIES"):
            with (
                self.subTest(constant=constant), patch(f"src.corpus.h2_practice_xlsx.{constant}", 1),
                self.assertRaises(ValueError),
            ):
                read_practice_workbook(data)

    def test_closes_workbook_after_failure(self) -> None:
        """Ошибка типа ячейки не оставляет открытый архив книги."""

        self.card["B14"] = 123
        data = _workbook_bytes(self.workbook)
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)

        with (
            patch("src.corpus.h2_practice_xlsx.load_workbook", return_value=workbook),
            patch.object(workbook, "close", wraps=workbook.close) as close,
            self.assertRaises(ValueError),
        ):
            read_practice_workbook(data)

        close.assert_called_once()


if __name__ == "__main__":
    unittest.main()
