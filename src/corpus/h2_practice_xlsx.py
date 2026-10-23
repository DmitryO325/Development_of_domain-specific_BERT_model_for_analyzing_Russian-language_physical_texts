"""Безопасное чтение текстовых ячеек Excel-бланков пробной разметки H2."""

from __future__ import annotations

import io
import re
import zipfile
import zlib
import xml.etree.ElementTree as ET

from dataclasses import dataclass
from pathlib import PurePosixPath

from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries
from openpyxl.utils.exceptions import InvalidFileException

MAX_INPUT_BYTES = 8 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 16 * 1024 * 1024
MAX_PART_BYTES = 4 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 256
MAX_WORKSHEETS = 32
MAX_ROWS = 200
MAX_COLUMNS = 8

_ALLOWED_PART = re.compile(
    r"(?:\[Content_Types\]\.xml|_rels/\.rels|docProps/[\w-]+\.xml|"
    r"xl/(?:workbook|styles|sharedStrings)\.xml|xl/_rels/workbook\.xml\.rels|"
    r"xl/theme/theme\d+\.xml|xl/worksheets/sheet\d+\.xml|"
    r"xl/worksheets/_rels/sheet\d+\.xml\.rels)"
)
_FORBIDDEN_ELEMENTS = {
    "f", "hyperlink", "hyperlinks", "comments", "externalReferences", "externalLink",
    "oleObject", "oleObjects", "controls", "control", "webExtensions", "connections",
}
_FORBIDDEN_RELATIONSHIPS = ("hyperlink", "comment", "external", "vba", "ole", "control", "package")


@dataclass(frozen=True)
class PracticeWorkbook:
    """Названия листов и точные непустые строки без вычисления формул."""

    sheets: tuple[str, ...]
    cells: dict[str, dict[str, str | None]]


def _local_name(tag: str) -> str:
    """Отбросить XML-пространство имён при проверке запрещённых элементов."""

    return tag.rsplit("}", 1)[-1]


def _check_xml(payload: bytes, part_name: str) -> None:
    """Проверить XML до передачи библиотеке, включая размеры и активное содержимое."""

    # Удаление нулевых байтов позволяет обнаружить декларации и в UTF-16/UTF-32.
    if re.search(br"<!\s*(?:DOCTYPE|ENTITY)\b", payload.replace(b"\x00", b""), re.IGNORECASE):
        raise ValueError("В Excel-бланке запрещены DTD и XML-сущности")

    document = ET.fromstring(payload)

    if sum(_local_name(element.tag) == "sheet" for element in document.iter()) > MAX_WORKSHEETS:
        raise ValueError(f"В Excel-бланке должно быть не больше {MAX_WORKSHEETS} листов")

    coordinates: set[str] = set()
    previous_row = 0

    for element in document.iter():
        name = _local_name(element.tag)

        if name in _FORBIDDEN_ELEMENTS:
            raise ValueError(f"В Excel-бланке запрещены формулы, ссылки и активное содержимое: {name}")

        if name == "Relationship":
            relationship = element.get("Type", "").rsplit("/", 1)[-1].lower()
            target = element.get("Target", "")

            if (
                element.get("TargetMode", "").lower() == "external"
                or any(token in relationship for token in _FORBIDDEN_RELATIONSHIPS)
                or ":" in target
                or target.startswith("//")
            ):
                raise ValueError("В Excel-бланке запрещены внешние связи, ссылки и вложения")

        if name in {"Override", "Default"}:
            content_type = element.get("ContentType", "").lower()

            if any(token in content_type for token in ("macro", "vba", "ole", "activex")):
                raise ValueError("В Excel-бланке запрещены макросы и встроенные объекты")

        if name == "sheet" and element.get("state", "visible") != "visible":
            raise ValueError("В Excel-бланке не должно быть скрытых листов")

        if name == "definedName" and element.get("name") not in {"_xlnm.Print_Area", "_xlnm.Print_Titles"}:
            raise ValueError("В Excel-бланке запрещены дополнительные именованные выражения")

        if not part_name.startswith("xl/worksheets/") or part_name.endswith(".rels"):
            continue

        if name == "dimension":
            minimum_column, minimum_row, maximum_column, maximum_row = range_boundaries(element.get("ref", ""))

            if (
                minimum_column is None or minimum_row is None
                or maximum_column is None or maximum_row is None
                or not 1 <= minimum_row <= maximum_row <= MAX_ROWS
                or not 1 <= minimum_column <= maximum_column <= MAX_COLUMNS
            ):
                raise ValueError(f"Размер листа превышает {MAX_ROWS} строк × {MAX_COLUMNS} столбцов")

        if name == "row":
            row_number = int(element.get("r", "0"))

            if not previous_row < row_number <= MAX_ROWS:
                raise ValueError("В Excel-бланке нарушен порядок или допустимое число строк")

            previous_row = row_number

        if name == "c":
            coordinate = element.get("r", "")
            row_number, column_number = coordinate_to_tuple(coordinate)

            if (
                coordinate in coordinates or row_number != previous_row
                or not 1 <= row_number <= MAX_ROWS or not 1 <= column_number <= MAX_COLUMNS
            ):
                raise ValueError("В Excel-бланке недопустимая или повторяющаяся координата ячейки")

            coordinates.add(coordinate)


def _check_archive(data: bytes) -> None:
    """Ограничить распаковку и разрешить только обычные части текстовой книги XLSX."""

    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()

        if len(entries) > MAX_ARCHIVE_ENTRIES:
            raise ValueError("В Excel-бланке слишком много частей ZIP-архива")

        if sum(entry.file_size for entry in entries) > MAX_UNCOMPRESSED_BYTES:
            raise ValueError("Распакованный Excel-бланк превышает допустимый размер")

        names: set[str] = set()

        for entry in entries:
            name = entry.filename
            path = PurePosixPath(name)

            if (
                name.lower() in names or path.is_absolute() or "\\" in name
                or any(part in {"", ".", ".."} for part in name.split("/"))
                or _ALLOWED_PART.fullmatch(name) is None
            ):
                raise ValueError(f"Недопустимая или повторяющаяся часть Excel-бланка: {name}")

            if (
                entry.file_size > MAX_PART_BYTES or entry.flag_bits & 1
                or entry.compress_type not in {zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED}
            ):
                raise ValueError("Часть Excel-бланка слишком велика, зашифрована или имеет неподдерживаемое сжатие")

            names.add(name.lower())

            with archive.open(entry) as source:
                payload = source.read(MAX_PART_BYTES + 1)

            if len(payload) > MAX_PART_BYTES:
                raise ValueError("Часть Excel-бланка превышает допустимый размер")

            _check_xml(payload, name)


def read_practice_workbook(data: bytes) -> PracticeWorkbook:
    """Прочитать ограниченную книгу в памяти, не изменяя файл и не вычисляя формулы."""

    if not isinstance(data, bytes) or not data:
        raise ValueError("Нужны непустые байты Excel-бланка XLSX")

    if len(data) > MAX_INPUT_BYTES:
        raise ValueError("Excel-бланк превышает допустимый размер файла")

    try:
        _check_archive(data)
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)

        try:
            if not 1 <= len(workbook.sheetnames) <= MAX_WORKSHEETS:
                raise ValueError(f"В Excel-бланке должно быть от 1 до {MAX_WORKSHEETS} листов")

            cells: dict[str, dict[str, str | None]] = {}

            for worksheet in workbook.worksheets:
                if worksheet.sheet_state != "visible":
                    raise ValueError("В Excel-бланке не должно быть скрытых листов")

                worksheet.reset_dimensions()
                sheet_cells: dict[str, str | None] = {}

                for row in worksheet.iter_rows(max_row=MAX_ROWS, max_col=MAX_COLUMNS):
                    for cell in row:
                        if cell.data_type in {"f", "e"}:
                            raise ValueError(f"Лист {worksheet.title}, {cell.coordinate}: формулы и ошибки запрещены")

                        value = cell.value

                        if value is not None and not isinstance(value, str):
                            raise ValueError(f"Лист {worksheet.title}, {cell.coordinate}: ожидается текст, а не число или дата")

                        if value not in {None, ""}:
                            sheet_cells[cell.coordinate] = value

                cells[worksheet.title] = sheet_cells

            return PracticeWorkbook(sheets=tuple(workbook.sheetnames), cells=cells)

        finally:
            workbook.close()

    except (
        zipfile.BadZipFile, ET.ParseError, InvalidFileException, KeyError, IndexError,
        AttributeError, TypeError, ValueError, OSError, EOFError, RuntimeError, zlib.error,
    ) as exception:
        raise ValueError(f"Не удалось прочитать Excel-бланк: {exception}") from exception
